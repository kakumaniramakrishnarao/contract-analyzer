import io
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import config, llm
from app.main import app

FIXTURES = Path(__file__).parent / "fixtures"


class FakeOpenAI:
    """Stands in for the OpenAI client. Queue replies (dicts, raw strings or exceptions)."""

    def __init__(self):
        self.replies = []
        self.calls = []
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=self._create))

    def _create(self, **kwargs):
        self.calls.append(kwargs["messages"])
        reply = self.replies.pop(0)
        if isinstance(reply, Exception):
            raise reply
        content = reply if isinstance(reply, str) else json.dumps(reply)
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content))])


@pytest.fixture
def fake_llm(monkeypatch):
    fake = FakeOpenAI()
    monkeypatch.setattr(llm, "_get_client", lambda: fake)
    return fake


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path / "data")
    return TestClient(app, raise_server_exceptions=False)


def pdf_bytes(name="service_agreement.pdf"):
    return (FIXTURES / name).read_bytes()


def docx_bytes():
    doc = Document()
    doc.add_heading("Mutual Non-Disclosure Agreement", level=1)
    doc.add_paragraph("This NDA is entered into by Bluebird Labs Inc. and Harbor Analytics LLC.")
    doc.add_heading("4. Governing Law", level=2)
    doc.add_paragraph("This NDA is governed by the laws of the State of New York.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Breach penalty"
    table.rows[0].cells[1].text = "$50,000"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload(client, data, name="contract.pdf"):
    return client.post("/upload/", files={"file": (name, data)})


@pytest.fixture
def pdf_doc(client):
    res = upload(client, pdf_bytes())
    assert res.status_code == 200, res.text
    return res.json()["doc_id"]
